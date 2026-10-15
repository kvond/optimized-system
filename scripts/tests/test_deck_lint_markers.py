"""Preserve active marker behavior using synthetic, student-free PPTX files.

These minimal fixtures intentionally omit unrelated deck requirements. Assertions
cover marker counts and findings, not whole-deck readiness.
"""
import importlib.util
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.util import Inches


LINTER_PATH = Path(__file__).resolve().parents[1] / "deck_lint.py"
SPEC = importlib.util.spec_from_file_location("deck_lint", LINTER_PATH)
deck_lint = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(deck_lint)


class MarkerDiagnosticsTests(unittest.TestCase):
    def lint_texts(self, *texts):
        presentation = Presentation()
        for text in texts:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5))
            box.text = text
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic.pptx"
            presentation.save(path)
            result = deck_lint.lint_deck(str(path))
        self.assertEqual(result["error"], "")
        return result

    @staticmethod
    def broken_markers(result):
        return [finding for finding in result["findings"]
                if finding.code == "H-MARKER-BROKEN"]

    def test_canonical_what_if_is_counted_case_insensitively(self):
        for heading in ("What if?", "WHAT IF?", "what if?"):
            with self.subTest(heading=heading):
                result = self.lint_texts(heading + " Water is unavailable.")
                self.assertEqual(result["diagnostic_slides"], 1)
                self.assertEqual(result["markers_present"], "What-if")
                self.assertEqual(self.broken_markers(result), [])

    def test_ellipsis_remains_a_hard_finding_with_slide_and_repair(self):
        text = "What if… water is unavailable."
        result = self.lint_texts("Introduction", text)
        self.assertEqual(result["diagnostic_slides"], 0)
        self.assertEqual(result["markers_present"], "")
        findings = self.broken_markers(result)
        self.assertEqual(len(findings), 1)
        finding = findings[0]
        self.assertEqual(finding.tier, deck_lint.HARD)
        self.assertEqual(finding.slide, 2)
        self.assertIn(text, finding.message)
        self.assertIn("Restore the exact ASCII marker in the deck", finding.message)
        self.assertNotIn("grader", finding.message)

    def test_recognizable_variants_do_not_become_canonical_markers(self):
        for text in ("What-if water is unavailable.", "What if water is unavailable."):
            with self.subTest(text=text):
                self.assertEqual(deck_lint.classify_slide(text), "What if")
                result = self.lint_texts(text)
                self.assertEqual(result["diagnostic_slides"], 0)
                self.assertEqual(result["markers_present"], "")
                self.assertEqual(self.broken_markers(result), [])

    def test_teacher_and_reflection_slides_stay_excluded(self):
        for exclusion in ("Do not project", "Teacher navigation",
                          "Continuation question:", "Relates to me:", "Concept Bank"):
            for marker in ("What if?", "What if…"):
                with self.subTest(exclusion=exclusion, marker=marker):
                    result = self.lint_texts(exclusion + "\n" + marker)
                    self.assertEqual(result["diagnostic_slides"], 0)
                    self.assertEqual(result["markers_present"], "")
                    self.assertEqual(self.broken_markers(result), [])

    def test_other_markers_and_multiple_markers_keep_existing_counts(self):
        result = self.lint_texts(
            "Critical aspect: Water availability",
            "Pattern break",
            "Finish this sentence as a rule",
            "Getting Started\nWorking On It\nMastery",
            "What if? Water is unavailable.\nPattern break",
            "Discuss what if water is unavailable.",
        )
        self.assertEqual(result["diagnostic_slides"], 5)
        self.assertEqual(result["tier_slides"], 1)
        self.assertEqual(result["markers_present"],
                         "Critical aspect; Pattern break; Build-a-rule; What-if")
        self.assertEqual(self.broken_markers(result), [])


if __name__ == "__main__":
    unittest.main()
